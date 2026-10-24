#!/usr/bin/env python3
"""Build manuscript_final.docx: prose identical to manuscript_final.md;
known math expressions are converted to native Word OMML equations.
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt
from latex2mathml.converter import convert as latex_to_mathml
from docx_equation import mathml_to_omml

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "manuscript" / "build" / "manuscript_final.md"
OUT = ROOT / "manuscript" / "build" / "manuscript_final.docx"

# Ordered rules: (regex, replacement latex template). Longest first.
RULES = [
    (r"n_hat = argmin_\{\|\|n\|\|=1\} \|\|Q_obs n - b\|\|\^2",
     r"\hat{n} = \text{arg min}_{\|n\|=1}\; \|Q_{\text{obs}}\, n - b\|^{2}"),
    (r"eps_theta = arccos\(n_hat\^T N0\)",
     r"\varepsilon_{\theta} = \arccos(\hat{n}^{T} N_{0})"),
    (r"S_k\(t\) = exp\(-\(t/tau_k\)\^a_k\)",
     r"S_{k}(t) = e^{-\left(t/\tau_{k}\right)^{a_{k}}}"),
    (r"N0 = L_Earth\(t0\)/\|\|L_Earth\(t0\)\|\|",
     r"N_{0} = L_{\text{Earth}}(t_{0}) \,/\, \|L_{\text{Earth}}(t_{0})\|"),
    (r"Q = \[q_1\^T; \.\.\.; q_n\^T\]",
     r"Q = [q_{1}^{T};\; \ldots;\; q_{n}^{T}]"),
    (r"H_ijk = sign\(det\(q_i, q_j, q_k\)\)",
     r"H_{ijk} = \text{sign}(\det(q_{i}, q_{j}, q_{k}))"),
    (r"b_i = q_i\^T N0", r"b_{i} = q_{i}^{T} N_{0}"),
    (r"b = Q N0", r"b = Q\, N_{0}"),
    (r"G = Q Q\^T", r"G = Q\, Q^{T}"),
    (r"G_ij = cos\(theta_ij\)", r"G_{ij} = \cos(\theta_{ij})"),
    (r"det\(O\) = -1", r"\det(O) = -1"),
    (r"kappa\(Q\^T Q\)", r"\kappa(Q^{T} Q)"),
    (r"P\(eps_theta < 1 ?deg\)", r"P(\varepsilon_{\theta} < 1\ \text{deg})"),
    (r"P\(eps ?< ?1 ?deg\)", r"P(\varepsilon < 1\ \text{deg})"),
    (r"P\(<1 ?deg\)", r"P(\varepsilon < 1\ \text{deg})"),
    (r"log10\(t/yr\)", r"\log_{10}(t/\text{yr})"),
    (r"log det", r"\log\,\det"),
    (r"B\(S\)", r"B(S)"),
    # token = value constructs
    (r"(sigma_obs|p_mismatch|p_invalid|N_MC|n_total|n_mc|tau_k|n)\s*"
     r"(>=|<=|=|<|>)\s*([0-9][0-9.,;e^+-]*(?:\s*(?:yr|rad|deg))?)",
     "__TOKVAL__"),
    (r"\bt = 10\^(\d+) ?yr", r"t = 10^{\1}\ \text{yr}"),
    (r"\b([ntcP]) ?(>=|<=|=|<|>) ?([0-9][0-9.,;]*)", r"\1 \2 \3"),
    (r"\bL_Earth\(t0\)", r"L_{\text{Earth}}(t_{0})"),
    # scientific notation
    (r"(?<![\w.])(\d+)e(-?\d+)(?![\w.])", r"\1 \times 10^{\2}"),
    (r"10\^(\d+)", r"10^{\1}"),
    # standalone identifier tokens
    (r"\beps_theta\b", r"\varepsilon_{\theta}"),
    (r"\bsigma_obs\b", r"\sigma_{\text{obs}}"),
    (r"\bp_mismatch\b", r"p_{\text{mismatch}}"),
    (r"\bp_invalid\b", r"p_{\text{invalid}}"),
    (r"\bn_total\b", r"n_{\text{tot}}"),
    (r"\bN_MC\b", r"N_{\text{MC}}"),
    (r"\bn_mc\b", r"n_{\text{mc}}"),
    (r"\bL_Earth\b", r"L_{\text{Earth}}"),
    (r"\bQ_obs\b", r"Q_{\text{obs}}"),
    (r"\bn_hat\b", r"\hat{n}"),
    (r"\bS_k\b", r"S_{k}"),
    (r"\btau_k\b", r"\tau_{k}"),
    (r"\ba_k\b", r"a_{k}"),
    (r"\bG_ij\b", r"G_{ij}"),
    (r"\bH_ijk\b", r"H_{ijk}"),
    (r"\btheta_ij\b", r"\theta_{ij}"),
    (r"\bN0\b", r"N_{0}"),
    (r"\bt0\b", r"t_{0}"),
    (r"\bq_i\b", r"q_{i}"),
    (r"\bq_j\b", r"q_{j}"),
    (r"\bq_k\b", r"q_{k}"),
    (r"\bq_1\b", r"q_{1}"),
    (r"\bq_n\b", r"q_{n}"),
    (r"\bq1\b", r"q_{1}"),
    (r"~(?=[0-9])", r"\approx "),
    (r"\bb_i\b", r"b_{i}"),
    (r"\bchi\b", r"\chi"),
]
TOKMAP = {
    "sigma_obs": r"\sigma_{\text{obs}}",
    "p_mismatch": r"p_{\text{mismatch}}",
    "p_invalid": r"p_{\text{invalid}}",
    "N_MC": r"N_{\text{MC}}",
    "n_total": r"n_{\text{tot}}",
    "n_mc": r"n_{\text{mc}}",
    "tau_k": r"\tau_{k}",
    "n": r"n",
}
UNITMAP = {"yr": r"\ \text{yr}", "rad": r"\ \text{rad}",
           "deg": r"\ \text{deg}"}


def tokval(m):
    tok, op, val = m.group(1), m.group(2), m.group(3)
    val_lx = re.sub(r"(\d+)e(-?\d+)",
                    lambda m: rf"{m.group(1)}\times10^{{{m.group(2)}}}", val)
    val_lx = re.sub(r"10\^(\d+)",
                    lambda m: rf"10^{{{m.group(1)}}}", val_lx)
    for w, lx in UNITMAP.items():
        val_lx = re.sub(rf"\s*{w}\b", lambda m, lx=lx: lx, val_lx)
    op_lx = {"=": "=", "<": "<", ">": ">",
             ">=": r"\geq", "<=": r"\leq"}[op]
    return f"{TOKMAP[tok]} {op_lx} {val_lx.strip()}"


COMPILED = [(re.compile(p), r) for p, r in RULES]


def split_math(text):
    """Return list of ('text'|'math', content-or-latex)."""
    out, pos = [], 0
    n = len(text)
    while pos < n:
        best = None
        for pat, tmpl in COMPILED:
            mm = pat.search(text, pos)
            if mm and (best is None or mm.start() < best[0] or
                       (mm.start() == best[0] and
                        mm.end() > best[1])):
                best = (mm.start(), mm.end(), tmpl, mm)
        if best is None:
            out.append(("text", text[pos:]))
            break
        s, e, tmpl, mm = best
        if s > pos:
            out.append(("text", text[pos:s]))
        if tmpl == "__TOKVAL__":
            lx = tokval(mm)
        else:
            lx = tmpl
            for g in range(mm.re.groups, 0, -1):
                lx = lx.replace(f"\\{g}", mm.group(g) or "")
        out.append(("math", lx))
        pos = e
    return out


ITALIC = re.compile(r"\*([^*\n]+)\*")


def add_plain(p, seg):
    last = 0
    for m in ITALIC.finditer(seg):
        if m.start() > last:
            p.add_run(seg[last:m.start()])
        p.add_run(m.group(1)).italic = True
        last = m.end()
    if last < len(seg):
        p.add_run(seg[last:])


def add_segments(p, s, problems=None):
    """Write text/OMML segments of one source line into paragraph p."""
    for kind, seg in split_math(s):
        if kind == "text":
            add_plain(p, seg)
        else:
            try:
                omml = mathml_to_omml(latex_to_mathml(seg))
                p._element.append(omml)
            except Exception as exc:  # noqa: BLE001
                if problems is not None:
                    problems.append((seg, str(exc)))
                p.add_run(seg)


def main():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    problems = []
    for line in SRC.read_text().splitlines():
        s = line.strip()
        if not s:
            continue
        if s.startswith("# "):
            doc.add_heading(s[2:], 0)
            continue
        if s.startswith("## "):
            doc.add_heading(s[3:], 1)
            continue
        if s.startswith("### "):
            doc.add_heading(s[4:], 2)
            continue
        if re.match(r"^[-*] ", s):
            p = doc.add_paragraph(style="List Bullet")
            s = s[2:]
        else:
            p = doc.add_paragraph()
        add_segments(p, s, problems)
    doc.save(OUT)
    print("wrote", OUT)
    for seg, e in problems:
        print("FALLBACK", repr(seg), e)


if __name__ == "__main__":
    main()
