"""Build manuscript_inline.docx: built manuscript text with figures and
tables embedded inline at their caption entries (author-review layout;
submission still uses separate figure/table files per journal rules).
"""
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "manuscript" / "build"
FIG = ROOT / "figures"
TAB = ROOT / "tables"

FIG_FILES = {
    "Figure 1": "fig1_architecture.png",
    "Figure 2": "fig2_p_vs_time.png",
    "Figure 3": "fig3_redundancy.png",
    "Figure 4": "fig4_error_vs_n.png",
    "Figure 5": "fig5_geometry.png",
    "Figure 6": "fig6_pareto.png",
    "Figure 7": "fig7_heatmap.png",
}
TAB_FILES = {
    "Table 1": "table1_architectures.csv",
    "Table 2": "table2_parameters.csv",
    "Table 3": "table3_performance.csv",
    "Table 4": "table4_sensitivity.csv",
}


def main():
    text = (BUILD / "manuscript_final.md").read_text()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    for line in text.splitlines():
        stripped = line.strip()
        if line.startswith("# "):
            doc.add_heading(line[2:], 0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], 1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], 2)
        elif not stripped:
            continue
        else:
            m = re.match(r"^(Figure \d)\.\s*(.*)$", stripped)
            t = re.match(r"^(Table \d)\.\s*(.*)$", stripped)
            if m and m.group(1) in FIG_FILES:
                f = FIG / FIG_FILES[m.group(1)]
                if f.exists():
                    doc.add_picture(str(f), width=Inches(5.5))
                doc.add_paragraph(f"{m.group(1)}. {m.group(2)}")
            elif t and t.group(1) in TAB_FILES:
                doc.add_paragraph(f"{t.group(1)}. {t.group(2)}")
                df = pd.read_csv(TAB / TAB_FILES[t.group(1)])
                tbl = doc.add_table(rows=1, cols=len(df.columns))
                tbl.style = "Table Grid"
                for j, c in enumerate(df.columns):
                    tbl.rows[0].cells[j].text = str(c)
                for _, row in df.iterrows():
                    cells = tbl.add_row().cells
                    for j, v in enumerate(row):
                        cells[j].text = "" if pd.isna(v) else str(v)
            else:
                doc.add_paragraph(stripped)

    out = BUILD / "manuscript_inline.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
