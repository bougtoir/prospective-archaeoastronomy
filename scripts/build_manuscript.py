#!/usr/bin/env python3
"""Build the final manuscript: substitute {{key}} placeholders from
manuscript/values.json into manuscript/main.md, number references in order
of appearance, emit manuscript/build/manuscript_final.md and a .docx.
"""
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
MS = ROOT / "manuscript"
BUILD = MS / "build"
BUILD.mkdir(exist_ok=True)


def _field(entry, name):
    """Extract a braced BibTeX field, tolerating nested braces."""
    m = re.search(rf"{name}\s*=\s*{{", entry)
    if not m:
        return None
    i = m.end()
    depth = 1
    j = i
    while j < len(entry) and depth:
        if entry[j] == "{":
            depth += 1
        elif entry[j] == "}":
            depth -= 1
        j += 1
    return entry[i:j - 1]


_ACCENT = {"'": 0x0301, '`': 0x0300, '"': 0x0308, '^': 0x0302,
           '~': 0x0303, '=': 0x0304, '.': 0x0307, 'u': 0x0306,
           'v': 0x030C, 'H': 0x030B, 'c': 0x0327, 'k': 0x0328,
           'r': 0x030A, 'b': 0x0331, 'd': 0x0323}


def _bibtex_to_unicode(s):
    """Convert BibTeX accent escapes like {\'a} or \'a to Unicode, and
    drop grouping braces."""
    def repl(m):
        return _compose(m.group(2), _ACCENT[m.group(1)])
    s = re.sub(r"\\(['`\"\^~=\.uvHckrbd])\{?([A-Za-z])\}?", repl, s)
    return s.replace("{", "").replace("}", "")


def _compose(letter, combining):
    import unicodedata
    return unicodedata.normalize("NFC", letter + chr(combining))


def main():
    text = (MS / "main.md").read_text()
    values = json.loads((MS / "values.json").read_text())

    def sub(m):
        key = m.group(1)
        if key not in values:
            raise SystemExit(f"placeholder without value: {key}")
        val = values[key]
        return "NA" if val is None else str(val)

    text = re.sub(r"\{\{([^}]+)\}\}", sub, text)

    # Renumber citations [key] -> [n] in order of appearance, and emit the
    # reference list in that order (Elsevier numbered style with titles).
    bib = {}
    bibsrc = MS / "references.bib"
    for m in re.finditer(r"@\w+\{([^,]+),(.*?)\n?\}\s*(?=@|\Z)",
                         bibsrc.read_text(), re.S):
        bib[m.group(1)] = m.group(2)
    order = []

    def cite(m):
        keys = [k.strip() for k in m.group(1).split(",")]
        nums = []
        for k in keys:
            if k not in bib:
                raise SystemExit(f"citation key not in bibliography: {k}")
            if k not in order:
                order.append(k)
            nums.append(order.index(k) + 1)
        return "[" + ",".join(str(n) for n in sorted(nums)) + "]"

    text = re.sub(r"\[([a-zA-Z_][\w:.\-]*(?:\s*,\s*[\w:.\-]+)*)\]", cite, text)

    refs = []
    for i, k in enumerate(order, 1):
        entry = bib[k]
        title = _field(entry, "title")
        author = _field(entry, "author")
        year = _field(entry, "year")
        jrnl = _field(entry, "journal")
        booktitle = _field(entry, "booktitle")
        publisher = _field(entry, "publisher")
        doi = _field(entry, "doi")
        line = f"[{i}] {_bibtex_to_unicode(author) if author else ''}"
        if title:
            line += f", \"{_bibtex_to_unicode(title)}\""
        if jrnl:
            line += f", {_bibtex_to_unicode(jrnl)}"
        elif booktitle:
            line += f", in {_bibtex_to_unicode(booktitle)}"
        elif publisher:
            line += f", {_bibtex_to_unicode(publisher)}"
        if year:
            line += f" ({year})"
        line += f". https://doi.org/{doi}" if doi else "."
        refs.append(line)
    text += "\n\n## References\n\n" + "\n\n".join(refs) + "\n"
    (BUILD / "manuscript_final.md").write_text(text)

    # ---- docx ----
    try:
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        for line in text.splitlines():
            if line.startswith("# "):
                doc.add_heading(line[2:], 0)
            elif line.startswith("## "):
                doc.add_heading(line[3:], 1)
            elif line.startswith("### "):
                doc.add_heading(line[4:], 2)
            elif line.strip():
                doc.add_paragraph(line)
        doc.save(BUILD / "manuscript_final.docx")
    except ImportError:
        print("python-docx not available; .docx skipped")

    (BUILD / "reference_map.json").write_text(json.dumps(order, indent=1))
    print("built manuscript:", len(order), "references")


if __name__ == "__main__":
    main()
